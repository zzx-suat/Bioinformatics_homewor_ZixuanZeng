#!/usr/bin/env python3
"""
把 Week 6 报告 md 转成 docx，并把代码、运行日志、sessionInfo 作为附录嵌进正文。

Week 4 的扣分教训：R 代码不在正文里 (-3)、缺真实运行输出 (-2)、
审计表不在被评文档里 (-3)。所以这里的所有必交元素都放进同一份 docx。

用法:  py R/build_docx.py      （工作目录为 week6_project）
输出:  Week6_分析报告_曾梓轩.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJ = "." if os.path.isdir("docs") else ".."
SRC = os.path.join(PROJ, "docs", "Week6_分析报告.md")
OUT = os.path.join(PROJ, "Week6_分析报告_曾梓轩.docx")

INK = RGBColor(0x24, 0x1A, 0x38)
ACCENT = RGBColor(0x3E, 0x2A, 0x63)
CN = "Microsoft YaHei"
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
CJK = re.compile(r"[　-〿一-鿿＀-￯]")

APPENDIX_CODE = ["R/01_load_clean.R", "R/02_alpha_diversity.R",
                 "R/03_beta_diversity.R", "R/04_differential.R", "R/run_all.R"]


def set_font(run, latin=None, east=CN):
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.append(rf)
    if latin:
        rf.set(qn("w:ascii"), latin)
        rf.set(qn("w:hAnsi"), latin)
    rf.set(qn("w:eastAsia"), east)


def add_runs(par, text, size=10.5):
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            r = par.add_run(chunk[2:-2]); r.bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1])
            r.font.size = Pt(size - 1.2); r.font.color.rgb = ACCENT
            set_font(r, latin="Consolas", east="Consolas")
            continue
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1]); r.italic = True
        else:
            r = par.add_run(chunk)
        r.font.size = Pt(size)
        if CJK.search(chunk):
            set_font(r)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), fill)
    tcPr.append(el)


def code_lines(doc, lines, size=7.6):
    for ln in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
        r = p.add_run(ln.rstrip("\n").replace("\t", "    "))
        r.font.size = Pt(size)
        set_font(r, latin="Consolas", east=CN)


def heading(doc, text, level):
    if level == 1:
        p = doc.add_paragraph()
        r = p.add_run(text); r.bold = True
        r.font.size = Pt(19); r.font.color.rgb = INK
        set_font(r)
        p.paragraph_format.space_after = Pt(10)
    else:
        h = doc.add_heading(text, level=min(level - 1, 4))
        for r in h.runs:
            r.font.color.rgb = ACCENT if level == 2 else INK
            set_font(r)


def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i


def main():
    if not os.path.exists(SRC):
        sys.exit(f"missing {SRC}")
    md = open(SRC, encoding="utf-8").read().split("\n")

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.85)
        s.top_margin = s.bottom_margin = Inches(0.8)
    doc.styles["Normal"].font.size = Pt(10.5)

    i = 0; n_tab = n_fig = 0
    while i < len(md):
        st = md[i].strip()

        if st == "---":
            i += 1; continue

        if st.startswith("```"):
            i += 1; buf = []
            while i < len(md) and not md[i].strip().startswith("```"):
                buf.append(md[i]); i += 1
            i += 1
            code_lines(doc, buf, size=8.6)
            doc.add_paragraph()
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", st)
        if m:
            heading(doc, m.group(2), len(m.group(1))); i += 1; continue

        if st.startswith("|"):
            rows, i = parse_table(md, i)
            if rows:
                n_tab += 1
                ncol = len(rows[0])
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = row[ci] if ci < len(row) else ""
                        c = t.cell(ri, ci); c.text = ""
                        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
                        add_runs(p, cell, size=8.8)
                        if ri == 0:
                            shade(c, "EDE9F2")
                            for r in p.runs: r.bold = True
                doc.add_paragraph()
            continue

        fm = re.match(r"^图：\s*`figures/([A-Za-z0-9_]+\.png)`", st)
        if fm:
            path = os.path.join(PROJ, "figures", fm.group(1))
            if os.path.exists(path):
                n_fig += 1
                pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(path, width=Inches(6.4))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(cap, f"**图 {n_fig}**　`figures/{fm.group(1)}`", size=9)
            i += 1; continue

        if st.startswith(">"):
            buf = []
            while i < len(md) and md[i].strip().startswith(">"):
                buf.append(md[i].strip().lstrip(">").strip()); i += 1
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, " ".join(x for x in buf if x))
            continue

        bm = re.match(r"^[-*]\s+(.*)$", st)
        nm = re.match(r"^(\d+)\.\s+(.*)$", st)
        if bm:
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, bm.group(1)); i += 1; continue
        if nm:
            p = doc.add_paragraph(style="List Number"); add_runs(p, nm.group(2)); i += 1; continue

        if not st:
            i += 1; continue
        p = doc.add_paragraph(); add_runs(p, st); i += 1

    # ---- 附录：代码、运行日志、sessionInfo ----
    doc.add_page_break()
    heading(doc, "附录 A — 完整分析代码", 2)
    p = doc.add_paragraph()
    add_runs(p, "以下为全部分析脚本原文，未节选。按 `R/run_all.R` 的顺序依次执行。")
    for f in APPENDIX_CODE:
        path = os.path.join(PROJ, f)
        heading(doc, f, 3)
        with open(path, encoding="utf-8") as fh:
            code_lines(doc, fh.readlines())
        doc.add_paragraph()

    doc.add_page_break()
    heading(doc, "附录 B — 完整运行日志", 2)
    p = doc.add_paragraph()
    add_runs(p, "`Rscript R/run_all.R` 在本机实际运行的控制台原样输出，未经编辑。")
    with open(os.path.join(PROJ, "results", "run_log.txt"), encoding="utf-8-sig") as fh:
        code_lines(doc, fh.readlines(), size=7.4)

    doc.add_page_break()
    heading(doc, "附录 C — 运行环境 sessionInfo()", 2)
    with open(os.path.join(PROJ, "results", "session_info.txt"), encoding="utf-8") as fh:
        code_lines(doc, fh.readlines(), size=7.6)

    doc.save(OUT)
    d = Document(OUT)
    print(f"wrote {OUT}")
    print(f"  段落 {len(d.paragraphs)} | 表格 {len(d.tables)} | 图片 {len(d.inline_shapes)}")
    print(f"  正文表格 {n_tab} | 嵌入图 {n_fig} | 附录代码文件 {len(APPENDIX_CODE)}")


if __name__ == "__main__":
    main()
