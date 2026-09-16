#!/usr/bin/env python3
"""
Build a Word version of the AI analysis, matching the structure used for
Weeks 1-3 (Heading 1/2, List Bullet, CodeBlk, tables, figures).

This converts the AI-side document only. The hand-written part is authored
directly as a .docx by R/build_worksheet_docx.py and is never generated from
markdown - it exists to be filled in by hand.

Usage:  py R/build_docx.py
In:     docs/Week4_AI_Analysis.md  + figures/*.png
Out:    Week4_AI分析_曾梓轩.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJ = "." if os.path.isdir("docs") else ".."
SRC = os.path.join(PROJ, "docs", "Week4_AI_Analysis.md")
OUT = os.path.join(PROJ, "Week4_AI分析_曾梓轩.docx")

INK = RGBColor(0x24, 0x1A, 0x38)
ACCENT = RGBColor(0x3E, 0x2A, 0x63)
MUTED = RGBColor(0x6B, 0x64, 0x78)
WARN = RGBColor(0xB4, 0x34, 0x1F)

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")
CJK_RE = re.compile(r"[　-〿一-鿿＀-￯]")
CJK_FONT = "Microsoft YaHei"


def set_cjk(run):
    """python-docx sets only the Latin font; CJK glyphs need w:eastAsia too,
    otherwise Word substitutes and the Chinese can render as boxes."""
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), CJK_FONT)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def add_runs(par, text, base_size=10.5, color=None):
    """Render **bold**, `code` and *italic* inside one paragraph."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            r = par.add_run(chunk[2:-2]); r.bold = True
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(base_size - 1.2)
            r.font.color.rgb = ACCENT
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1]); r.italic = True
        else:
            r = par.add_run(chunk)
        r.font.size = Pt(base_size)
        if CJK_RE.search(chunk):
            set_cjk(r)
        if color is not None and r.font.color.rgb is None:
            r.font.color.rgb = color


def ensure_styles(doc):
    styles = doc.styles
    if "CodeBlk" not in [s.name for s in styles]:
        from docx.enum.style import WD_STYLE_TYPE
        st = styles.add_style("CodeBlk", WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = "Consolas"; st.font.size = Pt(9)
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.left_indent = Inches(0.25)
    n = styles["Normal"]
    n.font.name = "Calibri"; n.font.size = Pt(10.5)
    n.paragraph_format.space_after = Pt(6)


def parse_table(lines, i):
    """Consume a pipe table starting at lines[i]; return (rows, next_i)."""
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i


def main():
    if not os.path.exists(SRC):
        sys.exit(f"missing {SRC}")
    md = open(SRC, encoding="utf-8").read().split("\n")

    doc = Document()
    ensure_styles(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)

    i = 0
    n_tables = n_figs = 0
    while i < len(md):
        line = md[i].rstrip()
        stripped = line.strip()

        # --- horizontal rule -> bottom-bordered spacer -----------------------
        if stripped == "---":
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
            bot.set(qn("w:color"), "D8D3E0")
            bd.append(bot); pPr.append(bd)
            p.paragraph_format.space_after = Pt(8)
            i += 1; continue

        # --- code fence -------------------------------------------------------
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(md) and not md[i].strip().startswith("```"):
                buf.append(md[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph(b, style="CodeBlk")
            continue

        # --- headings ---------------------------------------------------------
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            if lvl == 1:
                p = doc.add_paragraph()
                r = p.add_run(txt); r.bold = True
                r.font.size = Pt(19); r.font.color.rgb = INK
                if CJK_RE.search(txt):
                    set_cjk(r)
                p.paragraph_format.space_after = Pt(10)
            else:
                h = doc.add_heading(txt, level=min(lvl - 1, 4))
                for r in h.runs:
                    r.font.color.rgb = ACCENT if lvl == 2 else INK
                    if CJK_RE.search(r.text):
                        set_cjk(r)
            i += 1; continue

        # --- table ------------------------------------------------------------
        if stripped.startswith("|"):
            rows, i = parse_table(md, i)
            if rows:
                n_tables += 1
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci >= len(rows[0]):
                            continue
                        c = t.cell(ri, ci)
                        c.text = ""
                        par = c.paragraphs[0]
                        par.paragraph_format.space_after = Pt(2)
                        add_runs(par, cell, base_size=9)
                        if ri == 0:
                            shade(c, "EDE9F2")
                            for r in par.runs:
                                r.bold = True
                doc.add_paragraph()
            continue

        # --- figure embed -----------------------------------------------------
        fm = re.match(r"^\*\*Figure:\*\*\s+`figures/([A-Za-z0-9_]+\.png)`", stripped)
        if fm:
            fname = fm.group(1)
            path = os.path.join(PROJ, "figures", fname)
            cap = doc.add_paragraph()
            add_runs(cap, f"**Figure {n_figs + 1}.** `figures/{fname}`", base_size=9.5)
            if os.path.exists(path):
                n_figs += 1
                pic = doc.add_paragraph()
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(path, width=Inches(6.4))
            else:
                doc.add_paragraph(f"[missing figure: {path}]")
            i += 1; continue

        # --- blockquote -------------------------------------------------------
        if stripped.startswith(">"):
            buf = []
            while i < len(md) and md[i].strip().startswith(">"):
                buf.append(md[i].strip().lstrip(">").strip()); i += 1
            text = " ".join(x for x in buf if x)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
            left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "6"); left.set(qn("w:color"), "3E2A63")
            bd.append(left); pPr.append(bd)
            add_runs(p, text, base_size=10.5)
            continue

        # --- bullet / numbered list -------------------------------------------
        bm = re.match(r"^[-*]\s+(.*)$", stripped)
        nm = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if bm:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, bm.group(1)); i += 1; continue
        if nm:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, nm.group(2)); i += 1; continue

        # --- blank / body -----------------------------------------------------
        if not stripped:
            i += 1; continue
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"  tables embedded : {n_tables}")
    print(f"  figures embedded: {n_figs}")
    d = Document(OUT)
    print(f"  paragraphs      : {len(d.paragraphs)}")
    print(f"  inline shapes   : {len(d.inline_shapes)}")


if __name__ == "__main__":
    main()
